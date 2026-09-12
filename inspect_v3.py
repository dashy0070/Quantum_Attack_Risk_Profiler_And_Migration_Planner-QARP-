import docx

doc = docx.Document("2_CS_13_Capstone Project_Quantum_Attack_Risk_Profiler_And_Migration_Planner(QARP)_2026_v3.docx")
print("Total paragraphs in v3:", len(doc.paragraphs))

for i in range(len(doc.paragraphs)):
    t = doc.paragraphs[i].text.strip()
    if any(k in t.lower() for k in ["table 10.1", "chapter 10", "chapter 11", "conclusion", "sizing gap", "433.3", "300.0", "460.0", "200.0"]):
        print(f"P[{i}]: {t}\n")
