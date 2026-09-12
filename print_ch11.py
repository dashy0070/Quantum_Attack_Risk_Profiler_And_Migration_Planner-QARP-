import docx

doc = docx.Document("2_CS_13_Capstone Project_Quantum_Attack_Risk_Profiler_And_Migration_Planner(QARP)_2026_v3.docx")

for i in range(495, 515):
    if i < len(doc.paragraphs):
        print(f"P[{i}]: {doc.paragraphs[i].text}")
