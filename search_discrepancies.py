import docx
import os
import glob

docx_files = glob.glob("*.docx") + glob.glob("backup/*.docx")
print("Found docx files:", len(docx_files))

for path in docx_files:
    try:
        doc = docx.Document(path)
        for i, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if any(k in text.lower() for k in ["433", "460", "300.0%", "200.0%", "classical unit"]):
                print(f"\n--- {path} (Para {i}) ---")
                print(text)
        for t_idx, t in enumerate(doc.tables):
            for r_idx, row in enumerate(t.rows):
                row_txt = " | ".join(c.text.strip() for c in row.cells)
                if any(k in row_txt.lower() for k in ["433", "460", "300.0%", "200.0%"]):
                    print(f"\n--- {path} (Table {t_idx}, Row {r_idx}) ---")
                    print(row_txt)
    except Exception as e:
        print(f"Error reading {path}: {e}")
